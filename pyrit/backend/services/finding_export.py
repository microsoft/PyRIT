# Copyright (c) Microsoft Corporation.
# Licensed under the MIT license.

"""Selected findings of an Operation as a Word or Markdown section to copy into a report."""

import asyncio
from dataclasses import dataclass
from datetime import UTC, datetime
from functools import partial
from io import BytesIO
from urllib.parse import quote, urlencode
from uuid import UUID

from docx import Document
from docx.document import Document as DocxDocument
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt, RGBColor
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from pydantic import AnyHttpUrl

from pyrit.backend.models.attacks import ScoreView
from pyrit.backend.models.operations import (
    EvidenceAvailability,
    FindingEvidenceItem,
    FindingExportFormat,
    FindingExportRequest,
)
from pyrit.backend.services.operation_service import OperationService
from pyrit.models import Finding, FindingSeverity, ScoreStatus
from pyrit.models.harm_category import HarmCategory
from pyrit.output.finding_report.markdown import MarkdownFindingReportPrinter
from pyrit.output.finding_report.report import (
    DETAIL_SEPARATOR,
    EVIDENCE_NOTE,
    NO_EVIDENCE,
    UNAVAILABLE_EVIDENCE,
    FindingReport,
    FindingReportColumn,
    FindingReportEntry,
    FindingReportEvidence,
    FindingReportField,
    plain_text,
)

_MARKDOWN_MEDIA_TYPE = "text/markdown; charset=utf-8"
_DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_LINK_COLOR = RGBColor(0x05, 0x63, 0xC1)
_HUMAN_SCORER = "ManualScorer"
_BOOLEAN_SCORES = {"true": "True", "false": "False"}
# Word 2013 and later; the template's older mode opens documents in Compatibility Mode.
_WORD_COMPATIBILITY_MODE = "15"
_WORD_TABLE_STYLE = "Table Grid"
_WORD_COLUMNS = ("Finding", "Assessment")
# Each column's share of the page's text width: the Finding account needs most of the room.
_WORD_COLUMN_SHARES = (0.75, 0.25)
_WORD_BULLET_STYLE = "List Bullet"
# The template's List Bullet text indent, so an entry's lines align under its bullet text.
_WORD_BULLET_TEXT_INDENT = Pt(18)


@dataclass(frozen=True)
class FindingExportFile:
    """A generated findings document and the name to download it as."""

    content: bytes
    media_type: str
    filename: str


async def export_findings_async(*, operation_id: UUID, request: FindingExportRequest) -> FindingExportFile:
    """
    Build one document of an Operation's selected findings, in the dashboard's listing order, with evidence references.

    Evidence appears as saved metadata and dashboard links, without transcripts or media.

    Returns:
        FindingExportFile: The complete document.

    Raises:
        OperationNotFoundError: If the Operation does not exist.
        FindingNotFoundError: If a selected finding is not in the Operation.
    """
    report = await _build_report_async(operation_id=operation_id, request=request)
    now = datetime.now(UTC)
    if request.format == FindingExportFormat.MARKDOWN:
        extension, media_type = "md", _MARKDOWN_MEDIA_TYPE
        content = (await MarkdownFindingReportPrinter().render_async(report)).encode("utf-8")
    else:
        extension, media_type = "docx", _DOCX_MEDIA_TYPE
        content = await asyncio.to_thread(partial(_render_docx, report=report, now=now))
    filename = _export_filename(operation_id=operation_id, extension=extension, now=now)
    return FindingExportFile(content=content, media_type=media_type, filename=filename)


def _export_filename(*, operation_id: UUID, extension: str, now: datetime) -> str:
    """
    Name a download after its Operation ID, with the time in labelled UTC.

    Browsers save under the dashboard's own name, which uses the Operation name and the user's local time instead.

    Returns:
        str: For example ``copyrit-findings-c1e5502b-91f6-4d65-b182-4b028c7a75f6-2026-10-22-1658-utc.docx``.
    """
    return f"copyrit-findings-{operation_id}-{now.astimezone(UTC):%Y-%m-%d-%H%M}-utc.{extension}"


async def _build_report_async(*, operation_id: UUID, request: FindingExportRequest) -> FindingReport:
    service = OperationService()
    operation = await service.get_async(operation_id)
    selected = await service.get_findings_with_evidence_async(
        operation_id=operation_id, finding_ids=request.finding_ids
    )
    report_findings = []
    for finding, evidence in selected:
        severity, fields = _project_finding_fields(finding)
        report_findings.append(
            FindingReportEntry(
                title=finding.title,
                severity=severity,
                fields=fields,
                evidence=tuple(
                    _project_evidence(item=item, operation_id=operation_id, dashboard_origin=request.dashboard_origin)
                    for item in evidence
                ),
            )
        )
    return FindingReport(operation_name=operation.name, findings=tuple(report_findings))


def _project_finding_fields(finding: Finding) -> tuple[FindingReportField, tuple[FindingReportField, ...]]:
    """
    Label a finding's severity and its populated fields, in report order, and place each in a Word column.

    Narrative text goes in the Finding column and classifications in the Assessment column.

    Returns:
        tuple[FindingReportField, tuple[FindingReportField, ...]]: The severity, then the populated fields.
    """
    severity = (
        finding.severity_other or "Other"
        if finding.severity == FindingSeverity.OTHER
        else finding.severity.value.capitalize()
    )
    harm_type = finding.harm_type_other if finding.harm_type == HarmCategory.OTHER else finding.harm_type
    fields = [
        FindingReportField(label="Harm-type", value=str(harm_type), column=FindingReportColumn.ASSESSMENT)
        if harm_type
        else None,
        FindingReportField(label="Description", value=finding.description) if finding.description else None,
    ]
    return FindingReportField(label="Severity", value=severity), tuple(field for field in fields if field)


def _project_evidence(
    *, item: FindingEvidenceItem, operation_id: UUID, dashboard_origin: AnyHttpUrl
) -> FindingReportEvidence:
    if item.outcome is None:
        # The attack result is gone, so its saved ID stands in for its objective.
        return FindingReportEvidence(title=str(item.item.attack_result_id), summary=(), details=(), link_url=None)
    objective = item.objective if item.objective and item.objective.strip() else "No objective"
    summary = (
        FindingReportField(label="Outcome", value=item.outcome.value.capitalize()),
        FindingReportField(label="Score", value=_score_text(item.score)),
    )
    details = []
    if item.score is not None and item.score.score_rationale:
        details.append(FindingReportField(label="Rationale", value=item.score.score_rationale))
    # As on the dashboard, an attack's Operation is named only when it isn't the document's own.
    if item.source_operation_id != operation_id:
        details.append(
            FindingReportField(label="From operation", value=item.source_operation_name or "Not in an operation")
        )
    link_url = None
    if item.availability == EvidenceAvailability.AVAILABLE and item.conversation_id is not None:
        link_url = _conversation_link(
            dashboard_origin=dashboard_origin, item=item, conversation_id=item.conversation_id
        )
    return FindingReportEvidence(title=objective, summary=summary, details=tuple(details), link_url=link_url)


def _score_text(score: ScoreView | None) -> str:
    if score is None:
        return "Not scored"
    if score.status == ScoreStatus.UNDETERMINED:
        value = "Undetermined"
    else:
        value = _BOOLEAN_SCORES.get(str(score.score_value).lower(), str(score.score_value))
    scorer = "human review" if score.scorer_type == _HUMAN_SCORER else score.scorer_type
    return f"{value} ({scorer})"


def _conversation_link(*, dashboard_origin: AnyHttpUrl, item: FindingEvidenceItem, conversation_id: str) -> str:
    """
    Build the dashboard route that opens evidence's conversation and returns to the evidence.

    Returns:
        str: The absolute link, matching the dashboard's ``attackConversationRoutePath``; both are tested against
            ``frontend/src/utils/attackConversationRoute.cases.json``.
    """
    path = (
        f"/attacks/{_encode_segment(str(item.item.attack_result_id))}/conversations/{_encode_segment(conversation_id)}"
    )
    query = {"scenarioResultId": str(item.scenario_result_id)} if item.scenario_result_id else {}
    query["findingEvidenceId"] = str(item.item.id)
    return f"{str(dashboard_origin).rstrip('/')}{path}?{urlencode(query)}"


def _encode_segment(value: str) -> str:
    # The same characters JavaScript's encodeURIComponent leaves unescaped.
    return quote(value, safe="!*'()")


def _render_docx(*, report: FindingReport, now: datetime) -> bytes:
    document = Document()
    title = plain_text(f"Findings - {report.operation_name}")
    _replace_template_metadata(document=document, title=title, now=now)
    document.add_heading(title, level=1)
    document.add_paragraph(EVIDENCE_NOTE)
    table = document.add_table(rows=1, cols=len(_WORD_COLUMNS))
    table.style = document.styles[_WORD_TABLE_STYLE]
    header = table.rows[0]
    for cell, column in zip(header.cells, _WORD_COLUMNS, strict=True):
        cell.paragraphs[0].add_run(column).bold = True
    for finding in report.findings:
        finding_cell, assessment_cell = table.add_row().cells
        finding_cell.paragraphs[0].add_run(plain_text(finding.title)).bold = True
        for field in finding.fields:
            if field.column == FindingReportColumn.FINDING:
                _write_word_field(paragraph=finding_cell.add_paragraph(), field=field)
        _add_word_evidence(cell=finding_cell, evidence=finding.evidence)
        _write_word_field(paragraph=assessment_cell.paragraphs[0], field=finding.severity)
        for field in finding.fields:
            if field.column == FindingReportColumn.ASSESSMENT:
                _write_word_field(paragraph=assessment_cell.add_paragraph(), field=field)
    _set_column_widths(document=document, table=table)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _replace_template_metadata(*, document: DocxDocument, title: str, now: datetime) -> None:
    """Describe the export rather than python-docx's template, and open it outside Word's Compatibility Mode."""
    for setting in document.settings.element.iter(qn("w:compatSetting")):
        if setting.get(qn("w:name")) == "compatibilityMode":
            setting.set(qn("w:val"), _WORD_COMPATIBILITY_MODE)
    properties = document.core_properties
    properties.title = title
    properties.author = properties.comments = properties.last_modified_by = ""
    properties.created = properties.modified = now.astimezone(UTC)
    properties.revision = 1
    # The template's sample preview image and "Microsoft Macintosh Word" application details don't describe this file.
    package_relationships = document.part.package.rels
    for relationship_id, relationship in list(package_relationships.items()):
        if relationship.reltype in (RELATIONSHIP_TYPE.EXTENDED_PROPERTIES, RELATIONSHIP_TYPE.THUMBNAIL):
            del package_relationships[relationship_id]


def _set_column_widths(*, document: DocxDocument, table: Table) -> None:
    """Give each column its share of the page's text width, rather than letting Word size columns by content."""
    section = document.sections[0]
    page_width, left_margin, right_margin = section.page_width, section.left_margin, section.right_margin
    if page_width is None or left_margin is None or right_margin is None:
        # Without a page size and margins there is no text width to divide, so Word sizes the columns.
        return
    text_width = page_width - left_margin - right_margin
    table.autofit = False
    for column, share in zip(table.columns, _WORD_COLUMN_SHARES, strict=True):
        width = Emu(int(text_width * share))
        column.width = width
        # Word reads each cell's width and LibreOffice the column's, so both are set.
        for cell in column.cells:
            cell.width = width


def _add_word_evidence(*, cell: _Cell, evidence: tuple[FindingReportEvidence, ...]) -> None:
    if not evidence:
        _write_word_field(paragraph=cell.add_paragraph(), field=NO_EVIDENCE)
        return
    cell.add_paragraph().add_run(f"{NO_EVIDENCE.label}:").bold = True
    for entry in evidence:
        lead = cell.add_paragraph(style=_WORD_BULLET_STYLE)
        lead.paragraph_format.space_before = Pt(6)
        lead.paragraph_format.space_after = Pt(2)
        if entry.link_url is not None:
            _add_word_hyperlink(paragraph=lead, url=entry.link_url, text=plain_text(entry.title))
        else:
            lead.add_run(plain_text(entry.title)).bold = True
        if entry.summary:
            line = _indented_paragraph(cell)
            for index, field in enumerate(entry.summary):
                if index:
                    line.add_run(DETAIL_SEPARATOR)
                _write_word_field(paragraph=line, field=field)
        for field in entry.details:
            _write_word_field(paragraph=_indented_paragraph(cell), field=field)
        if entry.unavailable:
            _indented_paragraph(cell).add_run(UNAVAILABLE_EVIDENCE)


def _indented_paragraph(cell: _Cell) -> Paragraph:
    paragraph: Paragraph = cell.add_paragraph()
    paragraph.paragraph_format.left_indent = _WORD_BULLET_TEXT_INDENT
    paragraph.paragraph_format.space_after = Pt(2)
    return paragraph


def _write_word_field(*, paragraph: Paragraph, field: FindingReportField) -> None:
    paragraph.add_run(f"{plain_text(field.label)}: ").bold = True
    paragraph.add_run(plain_text(field.value))


def _add_word_hyperlink(*, paragraph: Paragraph, url: str, text: str) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    run = paragraph.add_run(text)
    run.font.underline = True
    run.font.color.rgb = _LINK_COLOR
    hyperlink.append(run._r)
    paragraph._p.append(hyperlink)
