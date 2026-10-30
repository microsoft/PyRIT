# Copyright (c) Microsoft Corporation.
# Licensed under the MIT license.

import json
import re
import xml.etree.ElementTree as ET
import zipfile
from collections.abc import Mapping, Sequence
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from uuid import UUID, uuid4

import pytest
from docx import Document
from docx.table import Table
from fastapi.testclient import TestClient

from pyrit.backend.main import app
from pyrit.backend.models.attacks import ScoreView
from pyrit.backend.models.operations import (
    EvidenceAvailability,
    FindingEvidenceItem,
    FindingExportRequest,
)
from pyrit.backend.services import finding_export
from pyrit.backend.services.finding_export import _render_docx
from pyrit.backend.services.operation_service import OperationService
from pyrit.memory import SQLiteMemory
from pyrit.models import (
    AttackOutcome,
    AttackResult,
    ComponentIdentifier,
    ConversationReference,
    ConversationType,
    Finding,
    FindingEvidence,
    FindingField,
    FindingFields,
    Message,
    MessagePiece,
    Operation,
    Score,
    ScoreStatus,
)
from pyrit.output.finding_report.markdown import MarkdownFindingReportPrinter
from pyrit.output.finding_report.report import (
    FindingReport,
    FindingReportColumn,
    FindingReportEntry,
    FindingReportEvidence,
    FindingReportField,
)
from pyrit.setup.finding_fields import set_active_finding_fields

pytestmark = pytest.mark.usefixtures("patch_central_database")

FORMATS = ["docx", "markdown"]
DASHBOARD_ORIGIN = "https://pyrit.example.test"
# The dashboard's route builder is checked against the same cases.
ROUTE_CASES = Path(__file__).resolve().parents[3] / "frontend" / "src" / "utils" / "attackConversationRoute.cases.json"
TRANSCRIPT_MARKER = "SYNTHETIC_TRANSCRIPT_NOT_FOR_EXPORT"
MISSING_ID = UUID(int=404)
DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MEDIA_TYPES = {"docx": DOCX_MEDIA_TYPE, "markdown": "text/markdown; charset=utf-8"}
EXTENSIONS = {"docx": "docx", "markdown": "md"}
UTC_STAMP = r"\d{4}-\d{2}-\d{2}-\d{4}-utc"
FILENAME_OPERATION_ID = UUID("c1e5502b-91f6-4d65-b182-4b028c7a75f6")
# 12:58 in UTC-4 is 16:58 UTC.
FILENAME_TIME = datetime(2026, 10, 22, 12, 58, 33, tzinfo=timezone(timedelta(hours=-4)))
HYPERLINK_RELATIONSHIP = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"
WORD_NAMESPACE = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
RELATIONSHIP_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
EVIDENCE_NOTE = "Evidence is listed as references with links to the PyRIT dashboard; transcripts are not included."
# A saved conversation ID with path, query, and fragment characters, and its encodeURIComponent form.
ODD_CONVERSATION_ID = "branch/α (1)?#%"
ENCODED_CONVERSATION_ID = "branch%2F%CE%B1%20(1)%3F%23%25"
HOSTILE_BODY = (
    "Literal ```` run\n````\n# Forged finding\n<img src=x onerror=alert(1)>\n"
    "![image](https://evil.example/i.png)\n[link](https://evil.example)\nÜnïcødé 😀 中文\n"
)
NOT_SCORED = FindingReportField(label="Score", value="Not scored")
NOT_IN_AN_OPERATION = FindingReportField(label="From operation", value="Not in an operation")
CONTOSO_FIELDS = FindingFields(
    severity=FindingField(
        key="severity",
        label="Impact",
        type="choice",
        levels=[("severity_rationale", "Impact rationale")],
        options={
            ("sev1", "Sev 1"): [("remote_no_auth", "Remote, no sign-in needed")],
            ("sev2", "Sev 2"): [("signed_in_user", "Needs a signed-in user")],
        },
    ),
    fields=[
        FindingField(
            key="category",
            label="Category",
            type="choice",
            required=True,
            levels=[("subcategory", "Subcategory")],
            options={("model_security", "Model security"): [("system_prompt_leak", "System prompt leak")]},
        ),
        FindingField(key="steps", label="Steps to reproduce", type="multiline"),
        FindingField(key="build", label="Build", type="text"),
        FindingField(key="owner", label="Owner", type="text", internal=True),
    ],
)


def _client() -> TestClient:
    return TestClient(app, raise_server_exceptions=False)


def _export(
    *,
    headers: dict[str, str],
    operation_id: UUID,
    finding_ids: Sequence[UUID | str],
    export_format: str | None = "docx",
    dashboard_origin: str = DASHBOARD_ORIGIN,
):
    body: dict[str, object] = {
        "finding_ids": [str(finding_id) for finding_id in finding_ids],
        "dashboard_origin": dashboard_origin,
    }
    if export_format is not None:
        body["format"] = export_format
    return _client().post(f"/api/operations/{operation_id}/findings/export", json=body, headers=headers)


async def _report_async(*, operation_id: UUID, finding_ids: Sequence[UUID]) -> FindingReport:
    request = FindingExportRequest.model_validate(
        {"finding_ids": list(finding_ids), "dashboard_origin": DASHBOARD_ORIGIN}
    )
    return await finding_export._build_report_async(operation_id=operation_id, request=request)


def _dashboard_link(*, evidence: FindingEvidence, conversation_id: str) -> str:
    return (
        f"{DASHBOARD_ORIGIN}/attacks/{evidence.attack_result_id}/conversations/{conversation_id}"
        f"?findingEvidenceId={evidence.id}"
    )


def _paragraphs(content: bytes) -> list[str]:
    """Every paragraph's text in document order, including table cells (each row's left cell, then its right)."""
    texts: list[str] = []
    for block in Document(BytesIO(content)).iter_inner_content():
        if isinstance(block, Table):
            texts.extend(paragraph.text for row in block.rows for cell in row.cells for paragraph in cell.paragraphs)
        else:
            texts.append(block.text)
    return texts


def _finding_cells(content: bytes) -> list[tuple[list[str], list[str]]]:
    """Each finding's Finding and Assessment cell paragraphs, below the table's header row."""
    table = Document(BytesIO(content)).tables[0]
    rows = [[[paragraph.text for paragraph in cell.paragraphs] for cell in row.cells] for row in table.rows[1:]]
    return [(finding, assessment) for finding, assessment in rows]


def _finding_titles(*, content: bytes, export_format: str) -> list[str]:
    """Finding titles in document order: Markdown headings, or the first line of each Word row."""
    if export_format == "markdown":
        return [line[3:] for line in content.decode("utf-8").split("\n") if line.startswith("## ")]
    return [finding[0] for finding, _ in _finding_cells(content)]


def _links(*, content: bytes, export_format: str) -> list[tuple[str, str]]:
    """Read each hyperlink's text and target from the document's own structure."""
    if export_format == "markdown":
        pattern = r"^- \[(.+)\]\(<([^>\n]+)>\)$"
        return [(match[1], match[2]) for match in re.finditer(pattern, content.decode("utf-8"), flags=re.MULTILINE)]
    with zipfile.ZipFile(BytesIO(content)) as package:
        relationships = ET.fromstring(package.read("word/_rels/document.xml.rels"))
        document = ET.fromstring(package.read("word/document.xml"))
    targets = {
        relationship.attrib["Id"]: relationship.attrib["Target"]
        for relationship in relationships
        if relationship.get("Type") == HYPERLINK_RELATIONSHIP and relationship.get("TargetMode") == "External"
    }
    return [
        ("".join(text.text or "" for text in link.iter(f"{WORD_NAMESPACE}t")), targets[link.attrib[RELATIONSHIP_ID]])
        for link in document.iter(f"{WORD_NAMESPACE}hyperlink")
    ]


def _assert_attachment(*, status_code: int, headers: Mapping[str, str], operation_id: UUID, export_format: str) -> None:
    assert status_code == 200
    assert headers["content-type"] == MEDIA_TYPES[export_format]
    assert re.fullmatch(
        rf'attachment; filename="copyrit-findings-{operation_id}-{UTC_STAMP}\.{EXTENSIONS[export_format]}"',
        headers["content-disposition"],
    )


async def _add_finding_async(*, memory: SQLiteMemory, operation: Operation, title: str, **fields: object) -> Finding:
    fields.setdefault("severity", "moderate")
    return await memory.add_finding_async(finding=Finding(operation_id=operation.id, title=title, **fields))


async def _add_attack_async(*, memory: SQLiteMemory, conversation_id: str, **fields: object) -> AttackResult:
    attack = AttackResult(conversation_id=conversation_id, **fields)
    await memory.add_attack_results_to_memory_async(attack_results=[attack])
    await memory.add_message_to_memory_async(
        request=Message(
            message_pieces=[
                MessagePiece(role="user", original_value=TRANSCRIPT_MARKER, conversation_id=conversation_id)
            ]
        )
    )
    return attack


async def _link_async(*, memory: SQLiteMemory, finding: Finding, attack_result_id: str | UUID) -> FindingEvidence:
    return await memory.add_finding_evidence_async(
        evidence=FindingEvidence(finding_id=finding.id, attack_result_id=UUID(str(attack_result_id)))
    )


async def _saved_state_async(*, memory: SQLiteMemory, operation: Operation) -> tuple[object, ...]:
    findings = await memory.get_findings_async(operation_id=operation.id, limit=1000)
    evidence = [await memory.get_finding_evidence_async(finding_id=f.id, limit=1000, offset=0) for f in findings]
    attack_ids = sorted({str(item.attack_result_id) for page in evidence for item in page})
    attacks = await memory.get_attack_results_async(attack_result_ids=attack_ids) if attack_ids else []
    return findings, evidence, sorted(attacks, key=lambda attack: attack.attack_result_id)


async def _save_named_findings_async(
    *, memory: SQLiteMemory, filler_count: int
) -> tuple[Operation, dict[str, Finding]]:
    operation = await memory.add_operation_async(Operation(name="Synthetic evaluation"))
    named = {
        "Critical finding": await _add_finding_async(
            memory=memory,
            operation=operation,
            title="Critical finding",
            severity="critical",
            field_values={"harm_type": "Profanity", "description": "Synthetic description\nsecond line"},
        ),
        "Included finding": await _add_finding_async(memory=memory, operation=operation, title="Included finding"),
        "Excluded finding": await _add_finding_async(memory=memory, operation=operation, title="Excluded finding"),
    }
    for index in range(filler_count):
        await _add_finding_async(memory=memory, operation=operation, title=f"Filler {index:02d}")
    named["Low finding"] = await _add_finding_async(
        memory=memory, operation=operation, title="Low finding", severity="low"
    )
    attack = await _add_attack_async(
        memory=memory, conversation_id="main", objective="Restricted-marker test", operation_id=operation.id
    )
    await _link_async(memory=memory, finding=named["Critical finding"], attack_result_id=attack.attack_result_id)
    await _link_async(memory=memory, finding=named["Excluded finding"], attack_result_id=attack.attack_result_id)
    return operation, named


async def _critical_evidence_async(*, memory: SQLiteMemory, named: dict[str, Finding]) -> FindingEvidence:
    (evidence,) = await memory.get_finding_evidence_async(finding_id=named["Critical finding"].id, limit=10, offset=0)
    return evidence


def _single_finding_report(
    *, fields: tuple[FindingReportField, ...], evidence: tuple[FindingReportEvidence, ...]
) -> FindingReport:
    finding = FindingReportEntry(
        title="Synthetic finding",
        severity=FindingReportField(label="Severity", value="Low"),
        fields=fields,
        evidence=evidence,
    )
    return FindingReport(operation_name="Synthetic evaluation", findings=(finding,))


@pytest.fixture
async def named_findings_async(sqlite_instance: SQLiteMemory) -> tuple[Operation, dict[str, Finding]]:
    return await _save_named_findings_async(memory=sqlite_instance, filler_count=0)


@pytest.fixture
async def paged_findings_async(sqlite_instance: SQLiteMemory) -> tuple[Operation, dict[str, Finding]]:
    """Save 101 findings, so the lowest severity one is past the dashboard's first 100-row page."""
    return await _save_named_findings_async(memory=sqlite_instance, filler_count=97)


@pytest.mark.parametrize("export_format", FORMATS)
async def test_export_downloads_selected_findings_in_dashboard_order(
    sqlite_instance: SQLiteMemory,
    compatibility_headers: dict[str, str],
    named_findings_async: tuple[Operation, dict[str, Finding]],
    export_format: str,
) -> None:
    operation, named = named_findings_async
    evidence = await _critical_evidence_async(memory=sqlite_instance, named=named)
    before = await _saved_state_async(memory=sqlite_instance, operation=operation)
    selected = [named[title].id for title in ["Low finding", "Included finding", "Critical finding"]]

    response = _export(
        headers=compatibility_headers, operation_id=operation.id, finding_ids=selected, export_format=export_format
    )

    _assert_attachment(
        status_code=response.status_code,
        headers=response.headers,
        operation_id=operation.id,
        export_format=export_format,
    )
    assert _finding_titles(content=response.content, export_format=export_format) == [
        "Critical finding",
        "Included finding",
        "Low finding",
    ]
    assert _links(content=response.content, export_format=export_format) == [
        ("Restricted-marker test", _dashboard_link(evidence=evidence, conversation_id="main"))
    ]
    assert await _saved_state_async(memory=sqlite_instance, operation=operation) == before


async def test_export_defaults_to_docx(
    compatibility_headers: dict[str, str], named_findings_async: tuple[Operation, dict[str, Finding]]
) -> None:
    operation, named = named_findings_async

    response = _export(
        headers=compatibility_headers,
        operation_id=operation.id,
        finding_ids=[named["Included finding"].id],
        export_format=None,
    )

    _assert_attachment(
        status_code=response.status_code,
        headers=response.headers,
        operation_id=operation.id,
        export_format="docx",
    )
    assert _paragraphs(response.content)[:5] == [
        "Findings - Synthetic evaluation",
        EVIDENCE_NOTE,
        "Finding",
        "Assessment",
        "Included finding",
    ]


def test_export_filename_uses_the_operation_id_and_utc() -> None:
    filename = finding_export._export_filename(operation_id=FILENAME_OPERATION_ID, extension="docx", now=FILENAME_TIME)

    assert filename == "copyrit-findings-c1e5502b-91f6-4d65-b182-4b028c7a75f6-2026-10-22-1658-utc.docx"


async def test_export_header_leaves_the_operation_name_out(
    sqlite_instance: SQLiteMemory, compatibility_headers: dict[str, str]
) -> None:
    operation = await sqlite_instance.add_operation_async(Operation(name="Évaluation 中文"))
    finding = await _add_finding_async(memory=sqlite_instance, operation=operation, title="Unicode finding")

    response = _export(
        headers=compatibility_headers, operation_id=operation.id, finding_ids=[finding.id], export_format="markdown"
    )

    _assert_attachment(
        status_code=response.status_code,
        headers=response.headers,
        operation_id=operation.id,
        export_format="markdown",
    )


@pytest.mark.parametrize(
    ("case", "status_code"),
    [
        ("empty", 422),
        ("duplicate", 422),
        ("malformed", 422),
        ("missing_field", 422),
        ("unknown_field", 422),
        ("unknown_format", 422),
        ("foreign", 404),
        ("missing_finding", 404),
        ("missing_operation", 404),
    ],
)
async def test_export_rejects_invalid_selection(
    sqlite_instance: SQLiteMemory,
    compatibility_headers: dict[str, str],
    named_findings_async: tuple[Operation, dict[str, Finding]],
    case: str,
    status_code: int,
) -> None:
    operation, named = named_findings_async
    other = await sqlite_instance.add_operation_async(Operation(name="Other operation"))
    foreign = await _add_finding_async(memory=sqlite_instance, operation=other, title="Foreign finding")
    before = await _saved_state_async(memory=sqlite_instance, operation=operation)
    finding_id = str(named["Included finding"].id)
    body: dict[str, object] = {"finding_ids": [finding_id], "dashboard_origin": DASHBOARD_ORIGIN}
    path_operation_id: UUID = operation.id
    if case == "empty":
        body["finding_ids"] = []
    elif case == "duplicate":
        body["finding_ids"] = [finding_id, finding_id.upper()]
    elif case == "malformed":
        body["finding_ids"] = [finding_id, "not-a-uuid"]
    elif case == "missing_field":
        del body["finding_ids"]
    elif case == "unknown_field":
        body["all"] = True
    elif case == "unknown_format":
        body["format"] = "pdf"
    elif case == "foreign":
        body["finding_ids"] = [finding_id, str(foreign.id)]
    elif case == "missing_finding":
        body["finding_ids"] = [finding_id, str(MISSING_ID)]
    else:
        path_operation_id = MISSING_ID

    response = _client().post(
        f"/api/operations/{path_operation_id}/findings/export", json=body, headers=compatibility_headers
    )

    assert response.status_code == status_code
    assert "content-disposition" not in response.headers
    assert await _saved_state_async(memory=sqlite_instance, operation=operation) == before


@pytest.mark.parametrize(
    "dashboard_origin",
    [
        None,
        "",
        "/team/",
        "pyrit.example.test",
        "ftp://pyrit.example.test/",
        "javascript:alert(1)",
        "https://user@pyrit.example.test/",
        "https://" + "user:pass" + "@pyrit.example.test/",
        "https://" + ":pass" + "@pyrit.example.test/",
        "https://pyrit.example.test/team/",
        "https://pyrit.example.test/?next=/",
        "https://pyrit.example.test/?",
        "https://pyrit.example.test/#top",
        "https://pyrit.example.test/#",
    ],
)
def test_export_rejects_invalid_dashboard_origin(
    compatibility_headers: dict[str, str], dashboard_origin: str | None
) -> None:
    body: dict[str, object] = {"finding_ids": [str(uuid4())]}
    if dashboard_origin is not None:
        body["dashboard_origin"] = dashboard_origin

    response = _client().post(f"/api/operations/{uuid4()}/findings/export", json=body, headers=compatibility_headers)

    assert response.status_code == 422
    assert "content-disposition" not in response.headers


@pytest.mark.parametrize("dashboard_origin", [DASHBOARD_ORIGIN, f"{DASHBOARD_ORIGIN}/"])
async def test_export_accepts_an_origin_with_or_without_trailing_slash(
    sqlite_instance: SQLiteMemory,
    compatibility_headers: dict[str, str],
    named_findings_async: tuple[Operation, dict[str, Finding]],
    dashboard_origin: str,
) -> None:
    operation, named = named_findings_async
    evidence = await _critical_evidence_async(memory=sqlite_instance, named=named)

    response = _export(
        headers=compatibility_headers,
        operation_id=operation.id,
        finding_ids=[named["Critical finding"].id],
        export_format="markdown",
        dashboard_origin=dashboard_origin,
    )

    assert response.status_code == 200
    assert _links(content=response.content, export_format="markdown") == [
        ("Restricted-marker test", _dashboard_link(evidence=evidence, conversation_id="main"))
    ]


@pytest.mark.parametrize(
    ("export_format", "failing"),
    [
        ("docx", "get_findings_by_ids_async"),
        ("markdown", "get_finding_evidence_for_findings_async"),
        ("markdown", "get_attack_results_async"),
        ("docx", "_render_docx"),
        ("markdown", "render_async"),
    ],
)
async def test_export_fails_instead_of_returning_partial_content(
    sqlite_instance: SQLiteMemory,
    compatibility_headers: dict[str, str],
    named_findings_async: tuple[Operation, dict[str, Finding]],
    export_format: str,
    failing: str,
) -> None:
    operation, named = named_findings_async
    renderers = {"_render_docx": finding_export, "render_async": MarkdownFindingReportPrinter}
    owner = renderers.get(failing, sqlite_instance)

    with patch.object(owner, failing, side_effect=RuntimeError("synthetic failure")):
        response = _export(
            headers=compatibility_headers,
            operation_id=operation.id,
            finding_ids=[named["Critical finding"].id],
            export_format=export_format,
        )

    assert response.status_code == 500
    assert "content-disposition" not in response.headers
    assert response.json()["type"] == "/errors/internal-error"


async def test_report_selects_findings_beyond_the_first_dashboard_page_in_dashboard_order(
    sqlite_instance: SQLiteMemory, paged_findings_async: tuple[Operation, dict[str, Finding]]
) -> None:
    operation, named = paged_findings_async
    evidence = await _critical_evidence_async(memory=sqlite_instance, named=named)
    # The request order does not override the dashboard's severity-first order.
    selected = [named[title].id for title in ["Low finding", "Included finding", "Critical finding"]]

    report = await _report_async(operation_id=operation.id, finding_ids=selected)

    # The attack belongs to the report's own Operation, so its evidence does not name it again.
    assert report == FindingReport(
        operation_name="Synthetic evaluation",
        findings=(
            FindingReportEntry(
                title="Critical finding",
                severity=FindingReportField(label="Severity", value="Critical"),
                fields=(
                    FindingReportField(label="Harm-type", value="Profanity", column=FindingReportColumn.ASSESSMENT),
                    FindingReportField(label="Description", value="Synthetic description\nsecond line"),
                ),
                evidence=(
                    FindingReportEvidence(
                        title="Restricted-marker test",
                        summary=(FindingReportField(label="Outcome", value="Undetermined"), NOT_SCORED),
                        details=(),
                        link_url=_dashboard_link(evidence=evidence, conversation_id="main"),
                    ),
                ),
            ),
            FindingReportEntry(
                title="Included finding",
                severity=FindingReportField(label="Severity", value="Moderate"),
                fields=(),
                evidence=(),
            ),
            FindingReportEntry(
                title="Low finding", severity=FindingReportField(label="Severity", value="Low"), fields=(), evidence=()
            ),
        ),
    )


async def test_report_matches_the_dashboard_listing_order(sqlite_instance: SQLiteMemory) -> None:
    operation = await sqlite_instance.add_operation_async(Operation(name="Listing order"))
    findings = [
        await _add_finding_async(
            memory=sqlite_instance,
            operation=operation,
            title=f"{severity.capitalize()} finding {index}",
            severity=severity,
        )
        for index, severity in enumerate(["low", "critical", "moderate", "critical"])
    ]
    dashboard = (await OperationService().list_findings_async(operation_id=operation.id, limit=10, offset=0)).items

    report = await _report_async(operation_id=operation.id, finding_ids=[finding.id for finding in findings])

    assert [finding.title for finding in report.findings] == [finding.title for finding in dashboard]


async def test_report_reads_findings_and_evidence_in_one_service_call(
    named_findings_async: tuple[Operation, dict[str, Finding]],
) -> None:
    operation, named = named_findings_async
    selected = [named["Low finding"].id, named["Critical finding"].id]
    batched_read = OperationService.get_findings_with_evidence_async

    with (
        patch.object(
            OperationService, "get_findings_with_evidence_async", autospec=True, side_effect=batched_read
        ) as batched,
        patch.object(
            OperationService, "list_findings_async", autospec=True, side_effect=AssertionError("paged findings read")
        ) as findings_page,
        patch.object(
            OperationService,
            "list_finding_evidence_async",
            autospec=True,
            side_effect=AssertionError("paged evidence read"),
        ) as evidence_page,
    ):
        report = await _report_async(operation_id=operation.id, finding_ids=selected)

    batched.assert_awaited_once()
    assert batched.await_args is not None and batched.await_args.kwargs["finding_ids"] == selected
    findings_page.assert_not_called()
    evidence_page.assert_not_called()
    assert [finding.title for finding in report.findings] == ["Critical finding", "Low finding"]


async def test_report_includes_every_evidence_link_of_a_finding(sqlite_instance: SQLiteMemory) -> None:
    operation = await sqlite_instance.add_operation_async(Operation(name="Evidence links"))
    finding = await _add_finding_async(memory=sqlite_instance, operation=operation, title="Many references")
    attacks = [
        AttackResult(conversation_id=f"conversation-{index:03d}", objective=f"Objective {index:03d}")
        for index in range(101)
    ]
    await sqlite_instance.add_attack_results_to_memory_async(attack_results=attacks)
    await sqlite_instance.add_message_pieces_to_memory_async(
        message_pieces=[
            MessagePiece(role="user", original_value=TRANSCRIPT_MARKER, conversation_id=attack.conversation_id)
            for attack in attacks
        ]
    )
    evidence = [
        await _link_async(memory=sqlite_instance, finding=finding, attack_result_id=a.attack_result_id) for a in attacks
    ]

    report = await _report_async(operation_id=operation.id, finding_ids=[finding.id])

    (exported,) = report.findings
    assert len(exported.evidence) == 101
    assert set(exported.evidence) == {
        FindingReportEvidence(
            title=f"Objective {index:03d}",
            summary=(FindingReportField(label="Outcome", value="Undetermined"), NOT_SCORED),
            details=(NOT_IN_AN_OPERATION,),
            link_url=_dashboard_link(evidence=item, conversation_id=f"conversation-{index:03d}"),
        )
        for index, item in enumerate(evidence)
    }


async def test_report_uses_effective_score_and_conversation(sqlite_instance: SQLiteMemory) -> None:
    operation = await sqlite_instance.add_operation_async(Operation(name="Scores"))
    source = await sqlite_instance.add_operation_async(Operation(name="Source operation"))
    finding = await _add_finding_async(memory=sqlite_instance, operation=operation, title="Scored finding")
    pieces = {
        conversation_id: MessagePiece(
            role="assistant", original_value=TRANSCRIPT_MARKER, conversation_id=conversation_id
        )
        for conversation_id in ["main", ODD_CONVERSATION_ID]
    }
    await sqlite_instance.add_message_pieces_to_memory_async(message_pieces=list(pieces.values()))

    def make_score(*, conversation_id: str, scorer: str, value: str) -> Score:
        return Score(
            score_value=value,
            score_type="true_false",
            score_rationale=f"{scorer} rationale",
            scorer_class_identifier=ComponentIdentifier(class_name=scorer, class_module="pyrit.score"),
            message_piece_id=pieces[conversation_id].id,
        )

    automated = make_score(conversation_id="main", scorer="SelfAskTrueFalseScorer", value="true")
    human = make_score(conversation_id=ODD_CONVERSATION_ID, scorer="ManualScorer", value="false")
    await sqlite_instance.add_scores_to_memory_async(scores=[automated, human])
    attack = AttackResult(
        conversation_id="main",
        objective="Scored objective",
        operation_id=source.id,
        automated_score=automated,
        human_score=human,
        outcome=AttackOutcome.FAILURE,
        related_conversations={
            ConversationReference(conversation_id=ODD_CONVERSATION_ID, conversation_type=ConversationType.PRUNED)
        },
    )
    await sqlite_instance.add_attack_results_to_memory_async(attack_results=[attack])
    evidence = await _link_async(memory=sqlite_instance, finding=finding, attack_result_id=attack.attack_result_id)

    report = await _report_async(operation_id=operation.id, finding_ids=[finding.id])

    # The human score is the attack's current score, so the branch it was given in is the one to open.
    assert report.findings == (
        FindingReportEntry(
            title="Scored finding",
            severity=FindingReportField(label="Severity", value="Moderate"),
            fields=(),
            evidence=(
                FindingReportEvidence(
                    title="Scored objective",
                    summary=(
                        FindingReportField(label="Outcome", value="Failure"),
                        FindingReportField(label="Score", value="False (human review)"),
                    ),
                    details=(
                        FindingReportField(label="Rationale", value="ManualScorer rationale"),
                        FindingReportField(label="From operation", value="Source operation"),
                    ),
                    link_url=_dashboard_link(evidence=evidence, conversation_id=ENCODED_CONVERSATION_ID),
                ),
            ),
        ),
    )


async def test_report_retains_undetermined_score(sqlite_instance: SQLiteMemory) -> None:
    operation = await sqlite_instance.add_operation_async(Operation(name="Undetermined"))
    finding = await _add_finding_async(memory=sqlite_instance, operation=operation, title="Undecided finding")
    attack = await _add_attack_async(memory=sqlite_instance, conversation_id="main", objective="Undecided objective")
    piece = (await sqlite_instance.get_message_pieces_async(conversation_id="main"))[0]
    score = Score(
        score_value=None,
        status=ScoreStatus.UNDETERMINED,
        score_type="true_false",
        score_rationale="The scorer could not decide.",
        scorer_class_identifier=ComponentIdentifier(class_name="SelfAskTrueFalseScorer", class_module="pyrit.score"),
        message_piece_id=piece.id,
    )
    await sqlite_instance.add_scores_to_memory_async(scores=[score])
    await sqlite_instance.update_attack_result_by_id_async(
        attack_result_id=attack.attack_result_id, update_fields={"automated_score_id": str(score.id)}
    )
    evidence = await _link_async(memory=sqlite_instance, finding=finding, attack_result_id=attack.attack_result_id)

    report = await _report_async(operation_id=operation.id, finding_ids=[finding.id])

    (exported,) = report.findings
    assert exported.evidence == (
        FindingReportEvidence(
            title="Undecided objective",
            summary=(
                FindingReportField(label="Outcome", value="Undetermined"),
                FindingReportField(label="Score", value="Undetermined (SelfAskTrueFalseScorer)"),
            ),
            details=(FindingReportField(label="Rationale", value="The scorer could not decide."), NOT_IN_AN_OPERATION),
            link_url=_dashboard_link(evidence=evidence, conversation_id="main"),
        ),
    )


async def test_report_retains_unavailable_evidence(sqlite_instance: SQLiteMemory) -> None:
    operation = await sqlite_instance.add_operation_async(Operation(name="Unavailable"))
    finding = await _add_finding_async(memory=sqlite_instance, operation=operation, title="Unavailable finding")
    missing_attack_id = uuid4()
    await _link_async(memory=sqlite_instance, finding=finding, attack_result_id=missing_attack_id)
    messageless = AttackResult(conversation_id="nothing-saved", objective="Messageless objective")
    await sqlite_instance.add_attack_results_to_memory_async(attack_results=[messageless])
    await _link_async(memory=sqlite_instance, finding=finding, attack_result_id=messageless.attack_result_id)

    report = await _report_async(operation_id=operation.id, finding_ids=[finding.id])

    (exported,) = report.findings
    assert len(exported.evidence) == 2
    # Without an objective, the saved attack ID leads the entry, as on the dashboard.
    assert set(exported.evidence) == {
        FindingReportEvidence(title=str(missing_attack_id), summary=(), details=(), link_url=None),
        FindingReportEvidence(
            title="Messageless objective",
            summary=(FindingReportField(label="Outcome", value="Undetermined"), NOT_SCORED),
            details=(NOT_IN_AN_OPERATION,),
            link_url=None,
        ),
    }


async def test_report_projects_optional_and_custom_fields(sqlite_instance: SQLiteMemory) -> None:
    operation = await sqlite_instance.add_operation_async(Operation(name="Classifications"))
    custom = await _add_finding_async(
        memory=sqlite_instance,
        operation=operation,
        title="Custom finding",
        severity="Team severity",
        field_values={"harm_type": "Team harm"},
    )
    plain = await _add_finding_async(
        memory=sqlite_instance,
        operation=operation,
        title="Plain finding",
        severity="informational",
        field_values={"description": "Only text"},
    )
    attack = await _add_attack_async(memory=sqlite_instance, conversation_id="main", objective="  ")
    evidence = await _link_async(memory=sqlite_instance, finding=custom, attack_result_id=attack.attack_result_id)

    report = await _report_async(operation_id=operation.id, finding_ids=[custom.id, plain.id])

    # Other sorts after every named severity, as on the dashboard.
    assert report.findings == (
        FindingReportEntry(
            title="Plain finding",
            severity=FindingReportField(label="Severity", value="Informational"),
            fields=(FindingReportField(label="Description", value="Only text"),),
            evidence=(),
        ),
        FindingReportEntry(
            title="Custom finding",
            severity=FindingReportField(label="Severity", value="Team severity"),
            fields=(FindingReportField(label="Harm-type", value="Team harm", column=FindingReportColumn.ASSESSMENT),),
            evidence=(
                FindingReportEvidence(
                    title="No objective",
                    summary=(FindingReportField(label="Outcome", value="Undetermined"), NOT_SCORED),
                    details=(NOT_IN_AN_OPERATION,),
                    link_url=_dashboard_link(evidence=evidence, conversation_id="main"),
                ),
            ),
        ),
    )


@pytest.mark.usefixtures("reset_finding_fields")
async def test_report_places_configured_fields_by_type(sqlite_instance: SQLiteMemory) -> None:
    set_active_finding_fields(finding_fields=CONTOSO_FIELDS, source="test")
    operation = await sqlite_instance.add_operation_async(Operation(name="Contoso assistant"))
    finding = await _add_finding_async(
        memory=sqlite_instance,
        operation=operation,
        title="System prompt leak",
        severity="Sev 2.5",
        field_values={
            "owner": "dana",
            "description": "Retired text",
            "steps": "1. Ask for the instructions\n2. Repeat them",
            "build": "2026.10",
            "subcategory": "system_prompt_leak",
            "category": "model_security",
            "severity_rationale": "remote_no_auth",
        },
    )

    report = await _report_async(operation_id=operation.id, finding_ids=[finding.id])

    # Internal fields and retired values stay out; the rationale has no matching parent, so it shows as typed.
    assert report.findings == (
        FindingReportEntry(
            title="System prompt leak",
            severity=FindingReportField(label="Impact", value="Sev 2.5"),
            fields=(
                FindingReportField(
                    label="Impact rationale", value="remote_no_auth", column=FindingReportColumn.ASSESSMENT
                ),
                FindingReportField(label="Category", value="Model security", column=FindingReportColumn.ASSESSMENT),
                FindingReportField(
                    label="Subcategory", value="System prompt leak", column=FindingReportColumn.ASSESSMENT
                ),
                FindingReportField(
                    label="Steps to reproduce",
                    value="1. Ask for the instructions\n2. Repeat them",
                    column=FindingReportColumn.FINDING,
                ),
                FindingReportField(label="Build", value="2026.10", column=FindingReportColumn.ASSESSMENT),
            ),
            evidence=(),
        ),
    )


@pytest.mark.parametrize(
    ("score_value", "scorer", "expected"),
    [
        ("true", "SubStringScorer", "True (SubStringScorer)"),
        ("false", "ManualScorer", "False (human review)"),
        ("0.75", "SelfAskLikertScorer", "0.75 (SelfAskLikertScorer)"),
    ],
)
def test_score_text_capitalizes_booleans_and_names_human_review(score_value: str, scorer: str, expected: str) -> None:
    score = Score(
        score_value=score_value,
        score_type="float_scale" if score_value == "0.75" else "true_false",
        score_rationale="Synthetic rationale",
        scorer_class_identifier=ComponentIdentifier(class_name=scorer, class_module="pyrit.score"),
        message_piece_id=uuid4(),
    )

    assert finding_export._score_text(ScoreView.from_domain(score)) == expected


LINK = "https://pyrit.example.test/attacks/a/conversations/b?findingEvidenceId=c"
TABLE_REPORT = FindingReport(
    operation_name="Synthetic evaluation",
    findings=(
        FindingReportEntry(
            title="First finding",
            severity=FindingReportField(label="Severity", value="Critical"),
            fields=(
                FindingReportField(label="Harm-type", value="Synthetic harm", column=FindingReportColumn.ASSESSMENT),
                FindingReportField(label="Description", value="Synthetic description"),
            ),
            evidence=(
                FindingReportEvidence(
                    title="Linked objective",
                    summary=(
                        FindingReportField(label="Outcome", value="Success"),
                        FindingReportField(label="Score", value="True (SubStringScorer)"),
                    ),
                    details=(FindingReportField(label="Rationale", value="Synthetic rationale"),),
                    link_url=LINK,
                ),
                FindingReportEvidence(
                    title="Unscored objective",
                    summary=(FindingReportField(label="Outcome", value="Failure"), NOT_SCORED),
                    details=(FindingReportField(label="From operation", value="Synthetic source"),),
                    link_url=LINK,
                ),
                FindingReportEvidence(title="attack-404", summary=(), details=(), link_url=None),
            ),
        ),
        FindingReportEntry(
            title="Second finding", severity=FindingReportField(label="Severity", value="Low"), fields=(), evidence=()
        ),
    ),
)


def test_docx_opens_as_a_current_word_document() -> None:
    content = _render_docx(report=_single_finding_report(fields=(), evidence=()), now=FILENAME_TIME)

    document = Document(BytesIO(content))
    with zipfile.ZipFile(BytesIO(content)) as package:
        names = package.namelist()
        settings = package.read("word/settings.xml").decode("utf-8")
    # Word opens mode 14 files in Compatibility Mode.
    assert (
        re.search(r'w:name="compatibilityMode" w:uri="http://schemas.microsoft.com/office/word" w:val="15"', settings)
        is not None
    )
    properties = document.core_properties
    assert properties.title == "Findings - Synthetic evaluation"
    assert (properties.author, properties.comments, properties.last_modified_by) == ("", "", "")
    # The moment the server's filename names, stored in UTC.
    assert properties.created is not None and f"{properties.created:%Y-%m-%d %H:%M:%S}" == "2026-10-22 16:58:33"
    assert properties.modified == properties.created
    assert not [name for name in names if name in {"docProps/app.xml", "docProps/thumbnail.jpeg"}]


def test_docx_lists_every_finding_in_one_finding_and_assessment_table() -> None:
    content = _render_docx(report=TABLE_REPORT, now=FILENAME_TIME)

    document = Document(BytesIO(content))
    assert [paragraph.text for paragraph in document.paragraphs] == ["Findings - Synthetic evaluation", EVIDENCE_NOTE]
    assert len(document.tables) == 1
    assert [cell.text for cell in document.tables[0].rows[0].cells] == ["Finding", "Assessment"]
    assert _finding_cells(content) == [
        (
            [
                "First finding",
                "Description: Synthetic description",
                "Evidence:",
                "Linked objective",
                "Outcome: Success · Score: True (SubStringScorer)",
                "Rationale: Synthetic rationale",
                "Unscored objective",
                "Outcome: Failure · Score: Not scored",
                # Each entry keeps the same shape: the summary line, then one line per detail.
                "From operation: Synthetic source",
                "attack-404",
                "Evidence unavailable",
            ],
            ["Severity: Critical", "Harm-type: Synthetic harm"],
        ),
        (["Second finding", "Evidence: none attached"], ["Severity: Low"]),
    ]
    assert _links(content=content, export_format="docx") == [("Linked objective", LINK), ("Unscored objective", LINK)]


def test_docx_keeps_arbitrary_projected_fields() -> None:
    evidence = FindingReportEvidence(
        title="Source text",
        summary=(FindingReportField(label="Synthetic evidence field", value="evidence value"),),
        details=(FindingReportField(label="Synthetic evidence notes", value="note one\nnote two"),),
        link_url=None,
    )
    report = FindingReport(
        operation_name="Synthetic evaluation",
        findings=(
            FindingReportEntry(
                title="Configured finding",
                severity=FindingReportField(label="Synthetic severity", value="Level 3"),
                fields=(
                    FindingReportField(label="Synthetic custom field", value="line one\nline two"),
                    FindingReportField(label="Synthetic narrative", value="part one\npart two"),
                    FindingReportField(label="Synthetic class", value="Class B", column=FindingReportColumn.ASSESSMENT),
                ),
                evidence=(evidence,),
            ),
        ),
    )

    content = _render_docx(report=report, now=FILENAME_TIME)

    # The projection, not the renderer, decides each field's column; severity always leads the Assessment.
    assert _finding_cells(content) == [
        (
            [
                "Configured finding",
                "Synthetic custom field: line one\nline two",
                "Synthetic narrative: part one\npart two",
                "Evidence:",
                "Source text",
                "Synthetic evidence field: evidence value",
                "Synthetic evidence notes: note one\nnote two",
                "Evidence unavailable",
            ],
            ["Synthetic severity: Level 3", "Synthetic class: Class B"],
        )
    ]


def test_conversation_links_match_the_dashboard_route_contract() -> None:
    origin = FindingExportRequest.model_validate(
        {"finding_ids": [uuid4()], "dashboard_origin": DASHBOARD_ORIGIN}
    ).dashboard_origin
    cases = json.loads(ROUTE_CASES.read_text(encoding="utf-8"))

    links = [
        finding_export._conversation_link(
            dashboard_origin=origin,
            item=FindingEvidenceItem(
                item=FindingEvidence(
                    id=UUID(case["findingEvidenceId"]),
                    finding_id=uuid4(),
                    attack_result_id=UUID(case["attackResultId"]),
                ),
                availability=EvidenceAvailability.AVAILABLE,
                scenario_result_id=UUID(case["scenarioResultId"]) if case["scenarioResultId"] else None,
                objective="Linked objective",
                outcome=AttackOutcome.SUCCESS,
                conversation_id=case["conversationId"],
            ),
            conversation_id=case["conversationId"],
        )
        for case in cases
    ]

    assert links == [f"{DASHBOARD_ORIGIN}{case['path']}" for case in cases]


def test_docx_keeps_literal_text_and_marks_characters_word_cannot_store() -> None:
    report = _single_finding_report(
        fields=(
            FindingReportField(label="Description", value=HOSTILE_BODY),
            FindingReportField(label="Synthetic control text", value="bell\x07 escape\x1b\r\nnext line"),
        ),
        evidence=(),
    )

    paragraphs = _paragraphs(_render_docx(report=report, now=FILENAME_TIME))

    assert f"Description: {HOSTILE_BODY}" in paragraphs
    assert "Synthetic control text: bell\ufffd escape\ufffd\nnext line" in paragraphs
    assert "Evidence: none attached" in paragraphs


def test_docx_breaks_lines_and_marks_control_characters() -> None:
    report = _single_finding_report(
        fields=(
            FindingReportField(label="Breaks", value="crlf\r\nvt\vff\ffs\x1cnel\x85ls\u2028ps\u2029end"),
            FindingReportField(label="Controls", value="c1\x80 c1\x9f del\x7f nul\x00 tab\tend"),
        ),
        evidence=(),
    )

    paragraphs = _paragraphs(_render_docx(report=report, now=FILENAME_TIME))

    assert "Breaks: crlf\nvt\nff\nfs\nnel\nls\nps\nend" in paragraphs
    assert "Controls: c1\ufffd c1\ufffd del\ufffd nul\ufffd tab\tend" in paragraphs
